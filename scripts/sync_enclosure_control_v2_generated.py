"""Attach the result-derived v2 supplement through existing builders."""
import shutil
from enclosure_control_v2_summary import ROOT, TITLE, summary, thesis_blocks, outline_text


def sync_pptx_outputs():
    from pptx import Presentation
    from build_thesis_pptx import add_title, add_footer, add_bullets, style_slide
    data = summary()
    if data is None:
        return
    hold = data['result']['aggregates']['holdout']
    bullets = (
        f"假設模型：{data['run_count']} 回合／{data['step_count']:,} 步",
        'fixed／PID＋FF／LQR／MPC holdout MAE：' + '/'.join(
            f"{hold[m]['tracking_mae_C']:.3f}" for m in ('fixed', 'pid', 'lqr', 'mpc')) + '°C',
        '相同候選回合預算；CPU 時間並不相等',
        '完整狀態與已知名義參數；參數變體並非實體跨設備確認',
        '超溫與數值失敗分開；實體機箱與 E8：NOT_EVALUATED',
    )
    for suffix in ('', '_30min'):
        path = ROOT / f'outputs/papers/thesis_presentation_zh{suffix}.pptx'
        deck = Presentation(path)
        matches = [slide for slide in deck.slides if any(
            TITLE in shape.text for shape in slide.shapes if shape.has_text_frame)]
        if len(matches) > 1:
            raise ValueError('Duplicate v2 slides')
        if matches:
            slide = matches[0]
            for shape in list(slide.shapes):
                shape._element.getparent().remove(shape._element)
        else:
            slide = deck.slides.add_slide(deck.slide_layouts[6])
        style_slide(slide)
        add_title(slide, TITLE, '2026-10-05；結果由 v2 evidence 產生')
        add_bullets(slide, .8, 1.6, 11.7, 5.2, bullets, level0_size=20)
        slide.notes_slide.notes_text_frame.text = '\n'.join(data[k] for k in ('text', 'overload', 'boundary'))
        add_footer(slide, len(deck.slides))
        deck.save(path)
        shutil.copy2(path, ROOT / f'docs/papers/thesis/thesis_presentation_zh{suffix}.pptx')
        outline = ROOT / f'docs/thesis/presentation_outline_zh{suffix}.md'
        content = outline.read_text()
        if TITLE not in content:
            outline.write_text(content + outline_text())
    notes = ROOT / 'docs/thesis/presentation_speaker_notes_zh_30min.md'
    if notes.exists() and TITLE not in notes.read_text():
        notes.write_text(notes.read_text() + outline_text())


def sync_docx_copy():
    shutil.copy2(ROOT / 'docs/papers/thesis/thesis_draft_zh.docx', ROOT / 'outputs/papers/thesis_draft_zh.docx')
