#!/usr/bin/env python3
"""Append the synchronized E12-E15 BMC evidence chain to thesis presentations."""

from __future__ import annotations

from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
TOKEN = "BMC E12–E15"
PPTX_PATHS = (
    ROOT / "outputs/papers/thesis_presentation_zh.pptx",
    ROOT / "outputs/papers/thesis_presentation_zh_30min.pptx",
)
OUTLINE_PATHS = (
    ROOT / "docs/thesis/presentation_outline_zh.md",
    ROOT / "docs/thesis/presentation_outline_zh_30min.md",
)
SPEAKER_NOTES = ROOT / "docs/thesis/presentation_speaker_notes_zh_30min.md"


def _append_outline(path: Path) -> None:
    content = path.read_text(encoding="utf-8")
    if TOKEN in content:
        return
    content += (
        "\n\n## BMC E12–E15：資料修正與同伺服器確認\n\n"
        "- E12：6 個 development files 未達 30 rows；final test 未開啟，NOT_EVALUATED。\n"
        "- E13：舊 parser／unit pipeline 的結果標記為 PARSER_INVALIDATED。\n"
        "- E14A/B：4,038 rows 完成來源與單位稽核；三檔 raw-unit regime 被正規化。\n"
        "- E14C retrospective：MAE 4.0882→1.8054°C，13/14 runs，95% CI [1.4271, 2.7939]°C。\n"
        "- E15：另 14 個先前未開啟檔案、3,112 rows，一次不重訓確認；H-ENC-08 支持，十項閘門全通過。\n"
        "- E15 MAE / RMSE / P95：1.6939 / 1.9672 / 3.5000 → 0.9744 / 1.2056 / 2.3888°C；12/14 runs 勝出。\n"
        "- Macro run MAE gain 0.6134°C；20,000 次 run-block bootstrap 95% CI [0.2721, 0.9831]°C。\n"
        "- 保留退步 runs：202308051757.csv（-0.5537°C）與 202310252230.csv（-0.3011°C）。\n"
        "- 只支持同伺服器 temporal/workload transfer；PC 機箱、NTC、跨伺服器、空間熱場與控制尚未驗證。機箱為目前應用方向，單房間仍屬整體論文範圍。\n"
    )
    path.write_text(content, encoding="utf-8")


def _append_speaker_notes() -> None:
    content = SPEAKER_NOTES.read_text(encoding="utf-8")
    if TOKEN in content:
        return
    content += (
        "\n\n## 補充投影片：BMC E12–E15 證據鏈\n\n"
        "E12 先因開發檔案不足而停止，E13 的輸出又暴露 parser 與單位問題。E14A/E14B 將來源 section 與單位制度分開修正後，E14C 的 frozen load-aware ridge 在已開啟的 14 個檔案上由 MAE 4.0882°C 降至 1.8054°C，13/14 runs 改善，區間下界大於零。這仍是回溯敏感度，不是獨立確認。E15 在檔案、模型內容 hash 與閘門固定後，對另 14 個先前未開啟檔案、3,112 rows 完成一次不重訓確認。MAE / RMSE / P95 由 1.6939 / 1.9672 / 3.5000 降至 0.9744 / 1.2056 / 2.3888°C；macro run MAE gain 為 0.6134°C，20,000 次 run-block bootstrap 95% CI 為 [0.2721, 0.9831]°C。12/14 runs 勝出、十項閘門全通過，H-ENC-08 支持；202308051757.csv 與 202310252230.csv 仍分別退步 0.5537 與 0.3011°C。證據僅支持同伺服器 temporal/workload transfer，不能外推至個人電腦機箱、NTC 硬體、跨伺服器、3-D 空間熱場或控制。機箱是目前應用推進方向，單房間仍是整體論文範圍。\n"
    )
    SPEAKER_NOTES.write_text(content, encoding="utf-8")


def _append_pptx(path: Path) -> None:
    if not path.exists():
        return
    from pptx import Presentation

    from build_thesis_pptx import add_bullets, add_footer, add_title, style_slide

    presentation = Presentation(path)
    existing = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    if TOKEN in existing:
        return
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    style_slide(slide)
    add_title(slide, "BMC E12–E15：修正與同伺服器確認")
    add_bullets(
        slide,
        0.9,
        1.55,
        11.5,
        5.25,
        [
            "E12：6 個 development files 未達 30 rows；final test 未開啟",
            "E13：舊 parser／unit pipeline 結果標記 PARSER_INVALIDATED",
            "E14A/B：4,038 rows 完成來源與單位稽核；三檔 raw-unit regime 被正規化",
            "E14C retrospective：MAE 4.0882 → 1.8054°C；13/14 runs；95% CI [1.4271, 2.7939]°C",
            "E15：另 14 檔、3,112 rows 不重訓確認；H-ENC-08 支持；12/14 runs 勝出",
            "E15 只支持同伺服器 temporal/workload transfer；E14C 仍屬回溯證據",
        ],
        level0_size=20,
    )
    add_footer(slide, len(presentation.slides))
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    style_slide(slide)
    add_title(slide, "E15：同伺服器 temporal/workload 確認")
    add_bullets(
        slide,
        0.9,
        1.55,
        11.5,
        5.25,
        [
            "14 個先前未開啟檔案、3,112 rows；frozen ridge 一次不重訓確認；十項閘門全通過",
            "MAE / RMSE / P95：1.6939 / 1.9672 / 3.5000 → 0.9744 / 1.2056 / 2.3888°C",
            "Macro run MAE gain 0.6134°C；95% CI [0.2721, 0.9831]°C；12/14 runs 勝出",
            "退步 runs：202308051757.csv（-0.5537°C）、202310252230.csv（-0.3011°C）",
            "只支持同伺服器轉移；PC 機箱、NTC、跨伺服器、空間熱場與控制尚未驗證",
            "目前應用方向為機箱；單房間仍屬整體論文範圍",
        ],
        level0_size=20,
    )
    add_footer(slide, len(presentation.slides))
    presentation.save(path)


def sync_pptx_outputs() -> None:
    for path in OUTLINE_PATHS:
        _append_outline(path)
    _append_speaker_notes()
    for path in PPTX_PATHS:
        _append_pptx(path)
