"""Evidence-derived supplement for every coupled artifact."""
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT/'openspec/changes/integrate-enclosure-session-20261012/artifacts'
TITLE = '機箱溫度流程整合與跨模型探索 2026-10-12'
LABELS = {'fixed':'固定風扇','pi':'PI','pid':'PID','lqr':'LQR','linear_mpc':'線性 MPC',
          'rank_h1':'單步排序','rank_h6':'六步排序','rank_h6_no_correction':'六步無校正'}
METHODS = tuple(LABELS)


def summary():
    if not (ARTIFACTS/'result.json').exists():
        return None
    result = json.loads((ARTIFACTS/'result.json').read_text())
    verification = json.loads((ARTIFACTS/'verification.json').read_text())
    if verification['status']!='PASS':
        raise ValueError('Verified evidence required for reporting')
    opened = result['aggregates']['open_loop']['holdout']
    closed = result['aggregates']['closed_loop']['holdout']
    text = (f"我把機箱溫度候選流程接入正式 Python 套件的獨立機箱介面，整合量測 CSV、有限參數辨識、稀疏估測、"
            f"LQR／線性 MPC 與候選排序。新的四組假設平台共 {result['run_count']} 回合、{result['step_count']:,} 步，"
            '以相同 900 秒校正辨識接觸熱導、風扇對流增益、共同熱容量倍率與風扇延遲；固定容量比例和其他係數仍是假設。'
            'holdout 包含不同比例的熱容量與模型沒有的熱源直接傳到空氣路徑，所有方法只取得金屬片、空氣、風扇、入口與功率。'
            '名義／辨識／稀疏校正估測 MAE 為 '
            +'/'.join(f"{opened[m]['source_mae_C']:.4f}" for m in ('nominal_physics','calibrated_physics','sparse_corrected'))
            +f"°C。平均稀疏校正改善 {100*result['decisions']['source_mae_reduction_fraction']:.2f}%，但只有 "
            f"{result['decisions']['source_wins_of_4']}/4 組改善，未達 3/4 判準，因此 H-ENC-51 不支持。")
    decisions = ('固定／PI／PID／LQR／線性 MPC／單步／六步／六步無校正追蹤 MAE 為 '
                 +'/'.join(f"{closed[m]['tracking_mae_C']:.4f}" for m in METHODS)
                 +f"°C。六步相對單步改善 {100*result['decisions']['horizon_tracking_reduction_fraction']:.3f}%，未達 5%，"
                 'H-ENC-52 不支持；PI 追蹤誤差低於六步排序，線性 MPC 的風扇能耗 proxy 較低但追蹤稍差。'
                 f"六步平均取樣超溫 {closed['rank_h6']['overtemp_s']:.2f} 秒／回合，全部閉迴路命令違例為零。"
                 '這些是不同閉迴路軌跡下的比較；不能把命令合規當成真實安全。')
    boundary = ('此次完成軟體流程、探索性模擬及實機實驗準備。相同有限校正是 commissioning-assisted synthetic transfer，'
                '並非零樣本或真實跨設備驗證；額外傳熱路徑只有一個指定反例，不證明任意方程家族泛化。'
                'LQR 經限幅後不再主張無限制最優性，MPC 使用當下線性化、六步預測與軟溫度代價。'
                '風扇 proxy 不是量測能耗。原室內濕度／照度和三維場維持原方法，未完成機箱三因子遷移；'
                '使用者於 2026-10-12 確認尚無設備或資料，實體辨識、NTC 精度與控制介入保留 NOT_EVALUATED；口頭報告尚無完成紀錄。')
    integrity = (f"全部 {verification['episodes']} 回合、{verification['rows']:,} 筆已重算指標，核對時間、"
                 '合成噪聲、原始感測警告、獨立熱動態與只讀觀測的流程回放；'
                 f"{verification['historical_files_unchanged']} 個歷史研究檔案未改動，一次性 E15 沒有重跑。"
                 '時間偏移、同步真值偏移、風扇端點、原始觀測與指標五種破壞案例均被拒絕。')
    return {'result':result,'text':text,'decisions':decisions,'boundary':boundary,'integrity':integrity}


def thesis_blocks():
    s = summary()
    return [] if s is None else [{'type':'heading','text':TITLE,'level':1}]+[
        {'type':'paragraph','text':s[key]} for key in ('text','decisions','boundary','integrity')]


def outline_text():
    s = summary()
    return '\n\n## '+TITLE+'\n\n'+'\n\n'.join(s[k] for k in ('text','decisions','boundary','integrity'))+'\n'


def ieee_text():
    s = summary(); r=s['result']; h=r['aggregates']['open_loop']['holdout']; c=r['aggregates']['closed_loop']['holdout']
    return (r'\subsection{Integrated Sparse Enclosure Session}'+'\n'
            r'A separate offline temperature session integrates bounded four-parameter commissioning, sparse state estimation, LQR, linear MPC and action ranking. '
            f"A fresh exploratory study uses {r['run_count']} episodes ({r['step_count']:,} samples), four assumed rigs, 900~s commissioning per rig, capacity-ratio mismatch and one unmodeled direct source--air path. "
            'All controllers receive only plate/air/fan/inlet/power observations. Nominal/calibrated/corrected holdout source MAE is '
            +r'/\allowbreak{}'.join(f"{h[m]['source_mae_C']:.4f}" for m in ('nominal_physics','calibrated_physics','sparse_corrected'))
            +r'~$^\circ$C; correction improves the macro mean but only 2/4 pairs, failing the prespecified 3/4 coverage gate (H-ENC-51 not supported). '
            'Fixed/PI/PID/LQR/linear-MPC/one-step/six-step/no-correction tracking MAE is '
            +r'/\allowbreak{}'.join(f"{c[m]['tracking_mae_C']:.4f}" for m in METHODS)
            +r'~$^\circ$C. Six-step ranking fails the 5\% improvement gate (H-ENC-52 not supported) and does not outperform PI. '
            f"Its sampled overtemperature averages {c['rank_h6']['overtemp_s']:.2f}~s; all commands meet shared bounds/slew. "
            r'Bounded LQR is not unconstrained optimal control; MPC uses a six-step affine linearization with soft thermal penalties. '
            r'Independent dynamics and input-only replay verify every sample; all prior study files remain unchanged. '
            r'This is commissioning-assisted assumed-model evidence, not hardware efficacy, arbitrary-family generality, measured energy or humidity/light/spatial transfer.'+'\n\n')


def sync_pptx_outputs():
    import shutil
    from pptx import Presentation
    from build_thesis_pptx import add_title,add_footer,add_bullets,style_slide
    s=summary()
    if s is None:
        return
    bullets = ('量測 CSV → 有限辨識 → 稀疏估測 → 可限制的控制輸出',
               '88 回合／10,560 步；熱容量與風扇延遲變體、額外傳熱路徑',
               '估測平均改善 21.19%，僅 2/4 組：覆蓋判準不支持',
               '六步改善僅 0.069%；未優於 PI，仍有取樣超溫',
               '新增稀疏觀測 LQR／線性 MPC；全部命令合規',
               '軟體與實機程序備妥；真實量測／介入待設備到位')
    for suffix in ('','_30min'):
        path=ROOT/f'outputs/papers/thesis_presentation_zh{suffix}.pptx'
        deck=Presentation(path)
        matches=[slide for slide in deck.slides if any(TITLE in shape.text for shape in slide.shapes if shape.has_text_frame)]
        if len(matches)>1:
            raise ValueError('Duplicate session supplement')
        slide=matches[0] if matches else deck.slides.add_slide(deck.slide_layouts[6])
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        style_slide(slide)
        add_title(slide,TITLE,'溫度軟體整合與探索性模擬；保留未達判準的結果')
        add_bullets(slide,.8,1.6,11.7,5.2,bullets,level0_size=20)
        slide.notes_slide.notes_text_frame.text='\n'.join(s[k] for k in ('text','decisions','boundary','integrity'))
        add_footer(slide,len(deck.slides))
        deck.save(path)
        shutil.copy2(path,ROOT/f'docs/papers/thesis/thesis_presentation_zh{suffix}.pptx')
        outline=ROOT/f'docs/thesis/presentation_outline_zh{suffix}.md'
        if TITLE not in outline.read_text():
            outline.write_text(outline.read_text()+outline_text())
    notes=ROOT/'docs/thesis/presentation_speaker_notes_zh_30min.md'
    if TITLE not in notes.read_text():
        notes.write_text(notes.read_text()+outline_text())


def sync_ieee_source():
    import re
    path=ROOT/'docs/papers/ieee/paper.tex'
    text=path.read_text()
    start='% BEGIN ENCLOSURE SESSION 20261012'
    end='% END ENCLOSURE SESSION 20261012'
    supplement=start+'\n'+ieee_text()+end+'\n\n'
    if start in text:
        text=re.sub(re.escape(start)+r'.*?'+re.escape(end)+r'\n*',lambda _:supplement,text,flags=re.S)
    else:
        anchor=r'\section{Discussion and Conclusion}'
        if anchor not in text:
            raise ValueError('IEEE insertion anchor missing')
        text=text.replace(anchor,supplement+anchor,1)
    path.write_text(text)
