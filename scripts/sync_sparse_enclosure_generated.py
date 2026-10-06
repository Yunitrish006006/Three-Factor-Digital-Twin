"""Use existing presentation builder conventions for an evidence-derived appendix."""
import shutil
from sparse_enclosure_summary import ROOT, TITLE, summary, outline_text, SUMMARY_KEYS


def sync_pptx_outputs():
    from pptx import Presentation
    from build_thesis_pptx import add_title, add_footer, add_bullets, style_slide
    s=summary()
    if s is None:
        return
    h=s['result']['aggregates']['open_loop']['holdout']
    bullets=(
        '物理主模型 → 金屬片／空氣稀疏校正 → 多步動作排序',
        f"108 回合；熱源估測 MAE {h['calibrated_physics']['source_mae_C']:.3f} → {h['sparse_corrected']['source_mae_C']:.3f}°C",
        '六步排序追蹤較單步改善 9.03%；仍未優於 PI',
        '指定假設平台；仍有超溫；實體辨識／介入未評估',
        '全速平衡 32.40°C 案例：散熱能力與控制誤差分開',
        '下週 LQR／Kalman 1960 原文介紹已備稿；尚未口頭報告',
    )
    for suffix in ('','_30min'):
        path=ROOT/f'outputs/papers/thesis_presentation_zh{suffix}.pptx'
        deck=Presentation(path)
        matches=[slide for slide in deck.slides if any(TITLE in shape.text for shape in slide.shapes if shape.has_text_frame)]
        if len(matches)>1:
            raise ValueError('Duplicate sparse candidate supplement')
        slide=matches[0] if matches else deck.slides.add_slide(deck.slide_layouts[6])
        if matches:
            for shape in list(slide.shapes):
                shape._element.getparent().remove(shape._element)
        style_slide(slide)
        add_title(slide,TITLE,'探索性溫度子問題；不擴張原方法的實測主張')
        add_bullets(slide,.8,1.6,11.7,5.2,bullets,level0_size=20)
        slide.notes_slide.notes_text_frame.text='\n'.join(s[k] for k in SUMMARY_KEYS)
        add_footer(slide,len(deck.slides))
        deck.save(path)
        shutil.copy2(path,ROOT/f'docs/papers/thesis/thesis_presentation_zh{suffix}.pptx')
        outline=ROOT/f'docs/thesis/presentation_outline_zh{suffix}.md'
        text=outline.read_text()
        if TITLE not in text:
            outline.write_text(text+outline_text())
    notes=ROOT/'docs/thesis/presentation_speaker_notes_zh_30min.md'
    if TITLE not in notes.read_text():
        notes.write_text(notes.read_text()+outline_text())
