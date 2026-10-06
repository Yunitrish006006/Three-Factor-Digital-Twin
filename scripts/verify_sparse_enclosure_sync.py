"""Verify candidate supplement content/copies; record visual QA separately."""
import hashlib
import json
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from sparse_enclosure_summary import ARTIFACTS, ROOT, TITLE, summary


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def office_text(path, member):
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read(member))
    return ''.join(node.text or '' for node in root.iter() if node.tag.endswith('}t'))


def main():
    s = summary()
    checks = []
    hashes = {}
    def require(condition, label):
        if not condition:
            raise ValueError(label)
        checks.append(label)
    for filename in ('thesis_draft_zh.docx', 'thesis_draft_zh.pdf',
                     'thesis_presentation_zh.pptx', 'thesis_presentation_zh_30min.pptx'):
        published = ROOT / 'docs/papers/thesis' / filename
        generated = ROOT / 'outputs/papers' / filename
        require(digest(published) == digest(generated), f'byte-equal copies: {filename}')
        hashes[str(published.relative_to(ROOT))] = digest(published)
    for relative in ('docs/thesis/thesis_draft_zh.md',
                     'docs/thesis/presentation_outline_zh.md',
                     'docs/thesis/presentation_outline_zh_30min.md',
                     'docs/thesis/presentation_speaker_notes_zh_30min.md'):
        content = (ROOT / relative).read_text()
        require(all(s[key] in content for key in ('text', 'boundary', 'lqr')),
                f'evidence-derived supplement: {relative}')
        hashes[relative] = digest(ROOT / relative)
    docx = ROOT / 'docs/papers/thesis/thesis_draft_zh.docx'
    content = office_text(docx, 'word/document.xml')
    require(all(s[key] in content for key in ('text', 'boundary', 'lqr')), 'DOCX supplement text')
    from pptx import Presentation
    slide_counts = {}
    for suffix in ('', '_30min'):
        relative = f'docs/papers/thesis/thesis_presentation_zh{suffix}.pptx'
        deck = Presentation(ROOT / relative)
        matches = [slide for slide in deck.slides if any(TITLE in shape.text for shape in slide.shapes if shape.has_text_frame)]
        require(len(matches) == 1, f'exactly one supplement slide: {relative}')
        slide = matches[0]
        notes = slide.notes_slide.notes_text_frame.text
        require(all(s[key] in notes for key in ('text', 'boundary', 'lqr')), f'slide notes preserve evidence/limits: {relative}')
        require(all(shape.left >= 0 and shape.top >= 0 and shape.left + shape.width <= deck.slide_width
                    and shape.top + shape.height <= deck.slide_height for shape in slide.shapes),
                f'supplement shape bounds: {relative}')
        slide_counts[relative] = len(deck.slides)
    tex = (ROOT / 'docs/papers/ieee/paper.tex').read_text()
    require(all(value in tex for value in ('0.9610/0.6580/0.3352', '108 episodes (19,440 steps)',
                                           'does not outperform PI', 'remain unevaluated')), 'IEEE text numbers and limits')
    pdf_pages = {}
    for relative, needles in (
        ('docs/papers/thesis/thesis_draft_zh.pdf', ('0.3352', '49.06', 'Kalman')),
        ('docs/papers/ieee/paper.pdf', ('0.9610/0.6580/0.3352', '19,440', 'does not outperform PI')),
    ):
        path = ROOT / relative
        text = subprocess.check_output(['pdftotext', str(path), '-'], text=True)
        compact = ''.join(text.split())
        require(all(''.join(value.split()) in compact for value in needles), f'PDF text numbers/limits: {relative}')
        pdf_pages[relative] = text.count('\f')
        hashes[relative] = digest(path)
    for relative in ('docs/papers/ieee/paper.tex', 'docs/papers/ieee/references.bib',
                     'docs/reports/professor_research_2026-10-06_zh.html',
                     'docs/research/lqr_kalman1960_reading_2026-10-06_zh.md',
                     'docs/reports/lqr_oral_script_2026-10-06_zh.md'):
        hashes[relative] = digest(ROOT / relative)
    report = {'status': 'PASS', 'checks': checks, 'sha256': hashes,
              'slide_counts': slide_counts, 'pdf_pages': pdf_pages,
              'observed_qa': {'html': '1440/820/390 px; anchors, details, Q/R keyboard, no horizontal body overflow or console errors; local links checked',
                              'pdf': 'thesis page 89; IEEE pages 7-8 inspected',
                              'office_full_render': 'PENDING: pdf2image missing from packaged renderer runtime',
                              'html_print': 'NOT_EVALUATED'},
              'scope': 'Content/copy/shape checks do not establish complete Office visual QA or hardware evidence.'}
    output = ARTIFACTS / 'artifact_sync_verification.json'
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(f'PASS: {len(checks)} synchronization checks; Office visual QA pending.')


if __name__ == '__main__':
    main()
