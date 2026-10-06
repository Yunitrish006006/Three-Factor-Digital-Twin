"""Append the MPC supplement through existing project builders."""
from pathlib import Path
import shutil

from enclosure_mpc_summary import ROOT, TITLE, outline_text, summary


def sync_docx_copy():
    source=ROOT/'docs/papers/thesis/thesis_draft_zh.docx'
    dest=ROOT/'outputs/papers/thesis_draft_zh.docx'
    dest.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(source,dest)


def sync_pptx_outputs():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from build_thesis_pptx import add_title, add_footer, add_bullets, style_slide
    data=summary()
    if data is None:
        return
    hold=data['result']['aggregates']['holdout']
    bullets=(
        '假設熱模型，36 回合與 4,320 控制步',
        'holdout MAE：固定 %.3f / PID %.3f / MPC %.3f°C' % tuple(hold[k]['tracking_mae_C'] for k in ('fixed','pid','mpc')),
        'nominal 無命令違例、超溫或求解回退',
        'MPC 風扇能耗 proxy 與 PWM 變化增加',
        '過載仍超溫並出現求解回退，溫度是軟限制',
        '實體機箱與 E8：NOT_EVALUATED，LQR：TODO',
    )
    for suffix in ('','_30min'):
        path=ROOT/f'outputs/papers/thesis_presentation_zh{suffix}.pptx'
        deck=Presentation(path)
        if not any(TITLE in shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame):
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            style_slide(slide)
            add_title(slide,TITLE,'2026-10-05 探索結果，參數為事前假設')
            add_bullets(slide,.8,1.6,11.7,5.2,bullets,level0_size=20)
            slide.notes_slide.notes_text_frame.text=data['text']+'\n'+data['boundary']
            add_footer(slide,len(deck.slides))
            deck.save(path)
        shutil.copy2(path,ROOT/f'docs/papers/thesis/thesis_presentation_zh{suffix}.pptx')
        outline=ROOT/f'docs/thesis/presentation_outline_zh{suffix}.md'
        content=outline.read_text()
        if TITLE not in content:
            outline.write_text(content+outline_text())
