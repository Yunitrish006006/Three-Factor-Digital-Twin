"""Content/copy verification separated from hash-bound observed visual QA."""
import hashlib
import json
from pathlib import Path
import subprocess
import xml.etree.ElementTree as ET
import zipfile

from artifact_visual_qa import visual_qa_status
from enclosure_session_summary import ROOT,ARTIFACTS,TITLE,summary,ieee_text


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def verify():
    data=summary();checks=[];hashes={}
    def require(condition,label):
        if not condition:
            raise ValueError(label)
        checks.append(label)
    for filename in ('thesis_draft_zh.docx','thesis_draft_zh.pdf','thesis_presentation_zh.pptx','thesis_presentation_zh_30min.pptx'):
        original=ROOT/'docs/papers/thesis'/filename
        require(sha(original)==sha(ROOT/'outputs/papers'/filename),'byte-equal copies: '+filename)
        hashes[str(original.relative_to(ROOT))]=sha(original)
    for relative in ('docs/thesis/thesis_draft_zh.md','docs/thesis/presentation_outline_zh.md',
                     'docs/thesis/presentation_outline_zh_30min.md','docs/thesis/presentation_speaker_notes_zh_30min.md'):
        text=(ROOT/relative).read_text()
        require(all(data[key] in text for key in ('text','decisions','boundary','integrity')),'Evidence/limits in '+relative)
        hashes[relative]=sha(ROOT/relative)
    with zipfile.ZipFile(ROOT/'docs/papers/thesis/thesis_draft_zh.docx') as archive:
        doc=ET.fromstring(archive.read('word/document.xml'))
        text=''.join(node.text or '' for node in doc.iter() if node.tag.endswith('}t'))
        require(all(data[key] in text for key in ('text','decisions','boundary','integrity')),'DOCX source synchronization')
    from pptx import Presentation
    slides={}
    for suffix in ('','_30min'):
        path=ROOT/f'docs/papers/thesis/thesis_presentation_zh{suffix}.pptx'
        deck=Presentation(path)
        match=[slide for slide in deck.slides if any(TITLE in shape.text for shape in slide.shapes if shape.has_text_frame)]
        require(len(match)==1,'One session slide '+suffix)
        require(all(data[key] in match[0].notes_slide.notes_text_frame.text for key in ('text','decisions','boundary','integrity')),
                'Evidence/limits in slide notes '+suffix)
        require(all(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=deck.slide_width
                    and shape.top+shape.height<=deck.slide_height for shape in match[0].shapes),'Slide bounds '+suffix)
        slides[str(path.relative_to(ROOT))]=len(deck.slides)
    require(ieee_text() in (ROOT/'docs/papers/ieee/paper.tex').read_text(),'IEEE evidence-derived section')
    pages={}
    for relative in ('docs/papers/thesis/thesis_draft_zh.pdf','docs/papers/ieee/paper.pdf'):
        text=subprocess.check_output(['pdftotext',str(ROOT/relative),'-'],text=True)
        compact=''.join(text.split())
        require(all(value in compact for value in ('1.0086/0.4681/0.3689','10,560','1.3755','H-ENC-51','H-ENC-52')),
                'PDF metrics/decisions '+relative)
        pages[relative]=text.count('\f')
        hashes[relative]=sha(ROOT/relative)
    html='docs/reports/professor_research_2026-10-12_zh.html'
    text=(ROOT/html).read_text()
    require(all(data[key] in text for key in ('text','decisions','boundary','integrity')),'HTML evidence synchronization')
    require(all(term in text for term in ('Kalman','p.104','p.114','900 秒','口頭報告尚無完成紀錄')),'LQR/source/scope in HTML')
    require((ROOT/'docs/reports/professor_oral_script_2026-10-12_zh.md').is_file(),'Oral script delivered')
    require((ROOT/'docs/research/enclosure_hardware_execution_2026-10-12_zh.md').is_file(),'Hardware execution preparation')
    observations=ARTIFACTS/'visual_qa_observations.json'
    qa=visual_qa_status(ROOT,observations,html,('desktop','tablet','mobile','anchors','interaction','console','local_links','offline_assets','print'))
    require(qa['status']=='PASS','Observed HTML and print QA must match final bytes')
    for relative in (*pages,'docs/reports/professor_research_2026-10-12_zh.pdf'):
        observed=visual_qa_status(ROOT,observations,relative,('inspected_pages',))
        require(observed['status']=='PASS','Observed PDF scope '+relative)
    for relative in (html,'docs/reports/professor_research_2026-10-12_zh.pdf','docs/papers/ieee/paper.tex',
                     'docs/papers/ieee/references.bib','scripts/build_thesis_docx.py','scripts/build_thesis_pptx.py'):
        hashes[relative]=sha(ROOT/relative)
    return {'status':'PASS','checks':checks,'sha256':hashes,'pdf_pages':pages,'slide_counts':slides,
            'html_observed_qa':qa,'software_and_experimental_preparation':'COMPLETE',
            'physical_measurements':'DEFERRED_NO_EQUIPMENT','oral_presentation':'NOT_RECORDED',
            'office_full_visual_render':'NOT_EVALUATED',
            'scope':'Content/copy integrity, final responsive HTML and full report print; thesis/IEEE only inspected changed PDF pages. Office rendering runtime unavailable.'}


if __name__=='__main__':
    result=verify()
    (ARTIFACTS/'delivery_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(f'PASS: {len(result["checks"])} delivery checks; software/preparation complete; Office visual QA unavailable')
