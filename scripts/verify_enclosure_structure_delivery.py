"""Coupled content/copy integrity and hash-bound observed HTML/PDF QA."""
import json
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from pptx import Presentation
from artifact_visual_qa import visual_qa_status
from enclosure_structure_summary import ROOT,ARTIFACTS,TITLE,summary,ieee_text,KEYS
from enclosure_session_summary import summary as previous_summary, TITLE as PREVIOUS_TITLE
from verify_enclosure_session_delivery import sha

PUBLIC_ARTIFACTS=ROOT/'openspec/changes/discover-public-thermal-scenes-20261012/artifacts'


def verify():
    data=summary(); previous=previous_summary(); checks=[]; hashes={}
    def require(condition,label):
        if not condition: raise ValueError(label)
        checks.append(label)
    for filename in ('thesis_draft_zh.docx','thesis_draft_zh.pdf','thesis_presentation_zh.pptx','thesis_presentation_zh_30min.pptx'):
        path=ROOT/'docs/papers/thesis'/filename
        require(sha(path)==sha(ROOT/'outputs/papers'/filename),'byte-equal copies: '+filename)
        hashes[str(path.relative_to(ROOT))]=sha(path)
    for relative in ('docs/thesis/thesis_draft_zh.md','docs/thesis/presentation_outline_zh.md',
                     'docs/thesis/presentation_outline_zh_30min.md','docs/thesis/presentation_speaker_notes_zh_30min.md'):
        text=(ROOT/relative).read_text()
        require(all(s[k] in text for s in (data,previous) for k in KEYS),'Both stages synchronized: '+relative)
        hashes[relative]=sha(ROOT/relative)
    with zipfile.ZipFile(ROOT/'docs/papers/thesis/thesis_draft_zh.docx') as archive:
        doc=ET.fromstring(archive.read('word/document.xml'))
        text=''.join(n.text or '' for n in doc.iter() if n.tag.endswith('}t'))
        require(all(s[k] in text for s in (data,previous) for k in KEYS),'DOCX content includes both stages')
    slides={}
    for suffix in ('','_30min'):
        path=ROOT/f'docs/papers/thesis/thesis_presentation_zh{suffix}.pptx'; deck=Presentation(path)
        for title,s in ((TITLE,data),(PREVIOUS_TITLE,previous)):
            match=[slide for slide in deck.slides if any(title in shape.text for shape in slide.shapes if shape.has_text_frame)]
            require(len(match)==1,'One synchronized slide: '+title+suffix)
            require(all(s[k] in match[0].notes_slide.notes_text_frame.text for k in KEYS),'Slide evidence and limits: '+title+suffix)
            require(all(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=deck.slide_width
                        and shape.top+shape.height<=deck.slide_height for shape in match[0].shapes),'Slide bounds: '+title+suffix)
        slides[str(path.relative_to(ROOT))]=len(deck.slides)
    require(ieee_text() in (ROOT/'docs/papers/ieee/paper.tex').read_text(),'IEEE evidence-derived supplement')
    pages={}
    for relative in ('docs/papers/thesis/thesis_draft_zh.pdf','docs/papers/ieee/paper.pdf'):
        text=subprocess.check_output(['pdftotext',str(ROOT/relative),'-'],text=True); compact=''.join(text.split())
        require(all(v in compact for v in ('1.0764','1.3468','11,520','H-ENC-71','H-ENC-72','H-ENC-51','H-ENC-52')),
                'Both stages PDF metrics and decisions: '+relative)
        pages[relative]=text.count('\f'); hashes[relative]=sha(ROOT/relative)
    html='docs/reports/professor_research_2026-10-12_zh.html'; text=(ROOT/html).read_text()
    require(all(s[k] in text for s in (data,previous) for k in KEYS),'HTML includes both stages and evidence limits')
    require(all(term in text for term in ('Kalman','p.104','p.114','1.8376','2.7155','口頭報告尚無完成紀錄')),'HTML LQR and adverse outcome coverage')
    for relative in ('docs/reports/professor_oral_script_2026-10-12_zh.md','docs/research/enclosure_hardware_execution_2026-10-12_zh.md'):
        require((ROOT/relative).is_file(),'Preparation delivered: '+relative)
    observations=(PUBLIC_ARTIFACTS if (PUBLIC_ARTIFACTS/'visual_qa_observations.json').exists() else ARTIFACTS)/'visual_qa_observations.json'
    if (PUBLIC_ARTIFACTS/'readiness.json').exists():
        require(all(term in text for term in ('id="online"','SimScale','HotSpot','TCLab','DARE 殘差','可穩定')),
                'Public scenes and practical LQR in HTML')
        for relative in ('docs/research/public_thermal_simulation_scenes_2026-10-12_zh.md',
                         'docs/research/lqr_practical_guide_2026-10-12_zh.md'):
            require((ROOT/relative).is_file(),'Source-linked preparation delivered: '+relative)
            hashes[relative]=sha(ROOT/relative)
    qa=visual_qa_status(ROOT,observations,html,('desktop','tablet','mobile','anchors','interaction','console','local_links','offline_assets','print'))
    require(qa['status']=='PASS','Final HTML observed QA matches current bytes')
    for relative in (*pages,'docs/reports/professor_research_2026-10-12_zh.pdf'):
        require(visual_qa_status(ROOT,observations,relative,('inspected_pages',))['status']=='PASS','Observed changed PDF pages: '+relative)
    for relative in (html,'docs/reports/professor_research_2026-10-12_zh.pdf','docs/papers/ieee/paper.tex','docs/papers/ieee/references.bib'):
        hashes[relative]=sha(ROOT/relative)
    return {'status':'PASS','checks':checks,'sha256':hashes,'pdf_pages':pages,'slide_counts':slides,
            'html_observed_qa':qa,'software_and_experimental_preparation':'COMPLETE',
            'physical_measurements':'DEFERRED_NO_EQUIPMENT','oral_presentation':'NOT_RECORDED',
            'office_visual_render':'OUTSIDE_HTML_REPORT_SCOPE',
            'scope':'Coupled source/copy content checks; current observed responsive HTML and full print. Unchanged native PDF visual scope reused only while exact bytes/evidence match prior observations.'}


if __name__=='__main__':
    report=verify()
    output=PUBLIC_ARTIFACTS if (PUBLIC_ARTIFACTS/'readiness.json').exists() else ARTIFACTS
    (output/'delivery_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(f'PASS: {len(report["checks"])} delivery checks; HTML and software/preparation complete')
