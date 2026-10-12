"""Evidence-derived second-stage supplement; never upgrades negative evidence."""
import json
from pathlib import Path
import re

ROOT=Path(__file__).resolve().parents[1]
ARTIFACTS=ROOT/'openspec/changes/select-enclosure-structure-20261012/artifacts'
TITLE='有限熱結構選擇與隱藏熱源反例 2026-10-12'
KEYS=('text','decisions','boundary','integrity')


def summary():
    if not (ARTIFACTS/'result.json').exists(): return None
    result=json.loads((ARTIFACTS/'result.json').read_text())
    audit=json.loads((ARTIFACTS/'verification.json').read_text())
    calibration=json.loads((ARTIFACTS/'calibration.json').read_text())
    if audit['status']!='PASS': raise ValueError('Verified structure evidence required')
    h=result['aggregates']['open_loop']['holdout']; c=result['aggregates']['closed_loop']['holdout']
    fits=[v for f in calibration.values() for v in f['candidates'].values()]
    cost=sum(v['wall_time_s'] for v in fits); calls=sum(v['residual_calls'] for v in fits)
    text=(f"我接著測試四參數共同容量模型 S 與七參數模型 F，後者分別辨識三個熱容量，並增加熱源到空氣的熱導。"
          '每個平台僅用前 120 筆校正資料辨識，再用後 60 筆可見感測的加權誤差選擇；F 至少改善 5% 且求解成功才採用，選擇後不重新擬合。'
          f"六個假設平台共 {result['run_count']} 回合、{result['step_count']:,} 筆新測試，校正合計 {len(fits)} 次擬合、"
          f"{calls} 次完整殘差呼叫、{cost:.3f} 秒。三個平台選 F，另三個保留 S；沒有用熱源真值選模型。")
    decisions=(f"新 holdout 的 S／選擇模型在稀疏校正後，熱源估測 MAE 為 {h['S_corrected']['source_mae_C']:.4f}／"
               f"{h['selected_corrected']['source_mae_C']:.4f}°C，平均誤差增加 {-100*result['decisions']['source_mae_reduction_fraction']:.2f}%，"
               f"僅 {result['decisions']['source_wins_of_6']}/6 組改善，H-ENC-71 不支持。"
               f"S／選擇模型的六步追蹤 MAE 為 {c['S_rank_h6']['tracking_mae_C']:.4f}／{c['selected_rank_h6']['tracking_mae_C']:.4f}°C，"
               f"增加 {-100*result['decisions']['tracking_reduction_fraction']:.2f}%，H-ENC-72 不支持；兩者平均取樣超溫皆為 {c['selected_rank_h6']['overtemp_s']:.2f} 秒／回合。"
               '全部閉迴路命令違例為零，仍不能推論溫度安全。')
    boundary=('可見金屬片／空氣／風扇擬合較好，仍不足以保證隱藏熱源更準。對流失配平台的 F 熱源容量達下界、空氣容量達上界，'
              '提醒我模型係數可能互相補償；局部 Jacobian 條件數不是全域可辨識性證明。'
              '這次只隔離模型結構與 PI／六步排序，沒有重做全部控制器排名；前一階段 LQR／MPC 比較與負面結果仍保留。'
              'F 保留為探索候選，未因校正擬合較好就升格為較佳方法。這是有限校正後的合成溫度實驗，沒有設備量測、實體控制或完整三因子遷移證據。')
    integrity=(f"我已重算全部 {audit['episodes']} 回合、{audit['rows']:,} 筆指標，核對另 {audit['calibration_rows']:,} 筆校正來源，"
               '檢查時序切分、選擇規則、獨立熱動態與只讀觀測回放；六種破壞案例均被拒絕。'
               f"{audit['historical_files_unchanged']} 個歷史證據檔案位元組未變，一次性 E15 與先前機箱實驗沒有重跑。")
    return {'result':result,'calibration':calibration,'text':text,'decisions':decisions,'boundary':boundary,'integrity':integrity}


def thesis_blocks():
    s=summary()
    return [] if s is None else [{'type':'heading','text':TITLE,'level':1}]+[{'type':'paragraph','text':s[k]} for k in KEYS]


def outline_text():
    s=summary(); return '\n\n## '+TITLE+'\n\n'+'\n\n'.join(s[k] for k in KEYS)+'\n'


def ieee_text():
    s=summary(); r=s['result']; h=r['aggregates']['open_loop']['holdout']; c=r['aggregates']['closed_loop']['holdout']
    return (r'\subsection{Chronological Thermal Structure Selection}'+'\n'
        'A separate exploratory extension compares a four-parameter common-capacity model S with a seven-parameter model F using independent capacities and direct source--air exchange. '
        'Per rig, only the first 120 commissioning observations are fitted with bounded nonlinear least squares~\\cite{scipyleastsquares}; the last 60 plate/air/fan observations select F only after a 5\\% weighted-error improvement. No refit or hidden-source selection occurs. '
        f"Six assumed rigs yield {r['run_count']} fresh episodes ({r['step_count']:,} samples). Three select F. Corrected S/selected holdout source MAE is "
        +r'/\allowbreak{}'.join(f"{h[m]['source_mae_C']:.4f}" for m in ('S_corrected','selected_corrected'))
        +r'~$^\circ$C: error increases 25.12\%, with only 2/6 pair wins (H-ENC-71 not supported). '
        'S/selected six-step tracking MAE is '+r'/\allowbreak{}'.join(f"{c[m]['tracking_mae_C']:.4f}" for m in ('S_rank_h6','selected_rank_h6'))
        +r'~$^\circ$C, a 2.95\% increase (H-ENC-72 not supported); both average 20~s sampled overtemperature. '
        r'Improved visible calibration fit does not ensure accurate hidden-source reconstruction. Capacity bounds under convection mismatch and local Jacobian conditioning are diagnostics, not global identifiability proofs. '
        r'All traces and selection decisions pass independent audits; earlier study bytes remain frozen. F remains an exploratory candidate, with no hardware or three-factor transfer claim.'+'\n\n')


def sync_ieee_source():
    path=ROOT/'docs/papers/ieee/paper.tex'; text=path.read_text()
    start='% BEGIN ENCLOSURE STRUCTURE 20261012'; end='% END ENCLOSURE STRUCTURE 20261012'
    supplement=start+'\n'+ieee_text()+end+'\n\n'
    if start in text: text=re.sub(re.escape(start)+r'.*?'+re.escape(end)+r'\n*',lambda _:supplement,text,flags=re.S)
    else:
        anchor=r'\section{Discussion and Conclusion}'
        if anchor not in text: raise ValueError('IEEE anchor missing')
        text=text.replace(anchor,supplement+anchor,1)
    path.write_text(text)


def sync_pptx_outputs():
    import shutil
    from pptx import Presentation
    from build_thesis_pptx import add_title,add_footer,add_bullets,style_slide
    s=summary()
    if s is None: return
    bullets=('S：共同容量；F：独立容量與熱源—空氣熱導',
             '前 120 筆辨識／後 60 筆選擇；不使用熱源真值',
             '96 回合／11,520 筆；三個平台選 F、三個保留 S',
             '隱藏熱源誤差增加 25.12%；僅 2/6 組改善',
             '六步追蹤誤差增加 2.95%；兩項判準皆不支持',
             '可見擬合較好不足以保证熱源準確；候選保留為探索')
    for suffix in ('','_30min'):
        path=ROOT/f'outputs/papers/thesis_presentation_zh{suffix}.pptx'; deck=Presentation(path)
        matches=[slide for slide in deck.slides if any(TITLE in shape.text for shape in slide.shapes if shape.has_text_frame)]
        if len(matches)>1: raise ValueError('Duplicate structure supplement')
        slide=matches[0] if matches else deck.slides.add_slide(deck.slide_layouts[6])
        for shape in list(slide.shapes): shape._element.getparent().remove(shape._element)
        style_slide(slide); add_title(slide,TITLE,'校正留出較好，新的隱藏熱源測試仍可能退步')
        add_bullets(slide,.8,1.6,11.7,5.2,[v.replace('独','獨').replace('保证','保證') for v in bullets],level0_size=20)
        slide.notes_slide.notes_text_frame.text='\n'.join(s[k] for k in KEYS)
        add_footer(slide,len(deck.slides)); deck.save(path)
        shutil.copy2(path,ROOT/f'docs/papers/thesis/thesis_presentation_zh{suffix}.pptx')
        outline=ROOT/f'docs/thesis/presentation_outline_zh{suffix}.md'
        if TITLE not in outline.read_text(): outline.write_text(outline.read_text()+outline_text())
    notes=ROOT/'docs/thesis/presentation_speaker_notes_zh_30min.md'
    if TITLE not in notes.read_text(): notes.write_text(notes.read_text()+outline_text())
